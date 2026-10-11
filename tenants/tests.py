"""
Aislamiento entre parqueos: un parqueo nunca ve, modifica ni elimina datos de otro.

Escenario: dos parqueos (Central y Norte), cada uno con tarifas, entradas,
suscripciones, baños y pagos. Placas «marcadoras» distintas por parqueo.
"""
import json
from datetime import timedelta
from decimal import Decimal
from io import BytesIO

import pytest
from django.utils.timezone import localdate, now
from openpyxl import load_workbook

from bathrooms.models import BathroomEntry, BathroomFee
from parking.models import Entry, Fee, PlatePolicy, Range
from payment.models import Payment, Subscription
from tenants.models import Parking, ParkingMember
from tenants.services import OPERATOR_GROUP, OWNER_GROUP, create_parking_account, ensure_role_groups

A_PLATE, B_PLATE = "PAAA111", "PBBB222"


def build_parking(parking, plate, fee_amount):
    """Datos completos de un parqueo; devuelve sus objetos principales."""
    fee = Fee.objects.create(parking=parking, name=f"Tarifa {parking.name}", default=True)
    Range.objects.create(fee=fee, start_minute=0, amount=Decimal(fee_amount))

    closed = Entry(parking=parking, plate=plate + "C", fee=fee,
                   entry_date_hour=now() - timedelta(hours=3),
                   departure_date_hour=now() - timedelta(hours=2))
    closed.save()
    active = Entry(parking=parking, plate=plate, fee=fee, entry_date_hour=now() - timedelta(minutes=30))
    active.save()

    policy = PlatePolicy.objects.create(parking=parking, plate=plate + "M", billing_type="MONTHLY", amount=Decimal("40"))
    bath_fee = BathroomFee.objects.create(parking=parking, name="Sanitario", amount=Decimal(fee_amount))
    BathroomEntry.objects.create(parking=parking, fee=bath_fee)

    sub = Subscription.objects.create(parking=parking, name=f"ParkOps {parking.name}", price=Decimal("25"), active=True)
    Payment.objects.create(subscription=sub, billing_date=localdate().replace(day=1) - timedelta(days=40), amount=Decimal("25"))

    return {"fee": fee, "closed": closed, "active": active, "policy": policy, "bath_fee": bath_fee}


@pytest.fixture
def world(parking, other_parking, django_user_model):
    a = build_parking(parking, A_PLATE, "1.00")
    b = build_parking(other_parking, B_PLATE, "7.00")

    groups = ensure_role_groups()
    owner_a = django_user_model.objects.create_user("duenoa", password="x")
    owner_a.groups.add(groups[OWNER_GROUP])
    ParkingMember.objects.create(user=owner_a, parking=parking)

    return {"A": parking, "B": other_parking, "a": a, "b": b, "owner_a": owner_a}


@pytest.fixture
def client_a(client, world):
    client.force_login(world["owner_a"])
    return client


def body(response):
    return response.content.decode()


# ---------------------------------------------------------------- Páginas

@pytest.mark.django_db
@pytest.mark.parametrize("url", [
    "/",
    "/parking/historial/",
    "/parking/suscripciones/",
    "/baños/entradas/",
    "/baños/tarifas/",
])
def test_paginas_solo_muestran_datos_del_parqueo(client_a, world, url):
    response = client_a.get(url)
    html = body(response)

    assert response.status_code == 200
    assert B_PLATE not in html
    assert "Tarifa Parqueo Norte" not in html


@pytest.mark.django_db
def test_historial_muestra_las_entradas_propias(client_a, world):
    from parking.utils import format_plate

    html = body(client_a.get("/parking/historial/"))
    assert format_plate(A_PLATE) in html


@pytest.mark.django_db
def test_dashboard_suma_solo_el_parqueo(client_a, world):
    context = client_a.get("/").context
    # Central: salida de hoy cobrada a $1; Norte cobraría $7
    assert context["total_daily_income"] == Decimal("1.00")
    assert context["daily_bathroom_income"] == Decimal("1.00")
    assert context["total_subscriptions_month_income"] == Decimal("40")


# ---------------------------------------------------------------- Acceso por ID (IDOR)

@pytest.mark.django_db
@pytest.mark.parametrize("method,url", [
    ("get", "/parking/salida/{active}"),
    ("post", "/parking/salida/{active}"),
    ("get", "/parking/editar/{active}/"),
    ("post", "/parking/editar/{active}/"),
    ("get", "/parking/ticket/?entry_id={active}"),
    ("get", "/parking/suscripciones/{policy}"),
    ("post", "/parking/suscripciones/estado/{policy}"),
    ("get", "/baños/tarifas/editar/{bath_fee}/"),
    ("post", "/baños/tarifas/estado/{bath_fee}/"),
    ("post", "/baños/entradas/registrar/{bath_fee}/"),
])
def test_no_se_accede_a_registros_de_otro_parqueo(client_a, world, method, url, django_user_model):
    world["owner_a"].user_permissions.add(
        *__import__("django.contrib.auth.models", fromlist=["Permission"]).Permission.objects.filter(
            codename__in=["change_entry", "change_platepolicy", "change_bathroomfee"]
        )
    )
    b = world["b"]
    url = url.format(active=b["active"].pk, policy=b["policy"].pk, bath_fee=b["bath_fee"].pk)

    response = getattr(client_a, method)(url)

    assert response.status_code == 404
    b["active"].refresh_from_db()
    b["policy"].refresh_from_db()
    b["bath_fee"].refresh_from_db()
    assert b["active"].state is True
    assert b["policy"].active is True
    assert b["bath_fee"].state is True
    assert BathroomEntry.objects.filter(parking=world["B"]).count() == 1


@pytest.mark.django_db
def test_superusuario_no_elimina_entrada_de_otro_parqueo_que_el_actual(admin_client, world):
    response = admin_client.post(f"/parking/eliminar/{world['b']['active'].pk}/")

    assert response.status_code == 404
    assert Entry.objects.filter(pk=world["b"]["active"].pk).exists()


@pytest.mark.django_db
def test_dueno_con_permiso_de_eliminar_no_borra_entradas_de_otro_parqueo(client, world):
    from django.contrib.auth.models import Permission

    world["owner_a"].user_permissions.add(Permission.objects.get(codename="delete_entry"))
    client.force_login(world["owner_a"])

    assert client.post(f"/parking/eliminar/{world['b']['active'].pk}/").status_code == 404
    assert Entry.objects.filter(pk=world["b"]["active"].pk).exists()

    assert client.post(f"/parking/eliminar/{world['a']['active'].pk}/").status_code == 302
    assert not Entry.objects.filter(pk=world["a"]["active"].pk).exists()


# ---------------------------------------------------------------- Búsqueda, formularios, voz

@pytest.mark.django_db
def test_placa_activa_en_otro_parqueo_se_registra_como_entrada(client_a, world):
    response = client_a.post("/parking/busqueda/", {"plate": B_PLATE})
    assert response.url == f"/parking/registro/{B_PLATE}"


@pytest.mark.django_db
def test_misma_placa_activa_en_dos_parqueos(world):
    entry = Entry(parking=world["A"], plate=B_PLATE, fee=world["a"]["fee"])
    entry.save()
    assert Entry.objects.filter(plate=B_PLATE, state=True).count() == 2


@pytest.mark.django_db
def test_misma_placa_con_suscripcion_en_dos_parqueos(world):
    PlatePolicy.objects.create(parking=world["A"], plate=world["b"]["policy"].plate, billing_type="DAILY", amount=Decimal("3"))
    assert PlatePolicy.objects.filter(plate=world["b"]["policy"].plate).count() == 2


@pytest.mark.django_db
def test_registro_solo_ofrece_y_acepta_tarifas_del_parqueo(client_a, world):
    form = client_a.get("/parking/registro/PNUEVO1").context["form"]
    assert list(form.fields["fee"].queryset) == [world["a"]["fee"]]

    response = client_a.post("/parking/registro/PNUEVO1", {"plate": "PNUEVO1", "fee": world["b"]["fee"].pk})
    assert response.status_code == 200
    assert "fee" in response.context["form"].errors
    assert not Entry.objects.filter(plate="PNUEVO1").exists()


@pytest.mark.django_db
def test_entrada_nueva_queda_en_el_parqueo_del_usuario(client_a, world):
    client_a.post("/parking/registro/PNUEVO2", {"plate": "PNUEVO2", "fee": world["a"]["fee"].pk})
    assert Entry.objects.get(plate="PNUEVO2").parking == world["A"]


@pytest.mark.django_db
def test_dictado_no_sugiere_placas_de_otro_parqueo(client_a, world):
    def dictate(text):
        return client_a.post(
            "/parking/busqueda/voz/", json.dumps({"alternatives": [text]}), content_type="application/json"
        ).json()

    # La placa activa del otro parqueo se trata como entrada nueva
    assert dictate(B_PLATE)["action"] == "entrada"

    # Una placa parecida no sugiere la del otro parqueo (PBBB228 ≈ PBBB222)
    assert not [c for c in dictate("PBBB228")["candidates"] if c["plate"].startswith("PBBB")]

    # Pero sí sugiere la propia (PAAA118 ≈ PAAA111)
    assert A_PLATE in [c["plate"] for c in dictate("PAAA118")["candidates"]]


# ---------------------------------------------------------------- Reportes

def excel_plates(response):
    sheet = load_workbook(BytesIO(response.content)).active
    return {row[0] for row in sheet.iter_rows(min_row=4, values_only=True) if row[0]}


@pytest.mark.django_db
@pytest.mark.parametrize("url", [
    "/parking/reporte/hoy/?date={d}&format=xlsx",
    "/parking/reporte/mes/?month_date={m}&format=xlsx",
    "/parking/reporte/periodo/?period_start_date={d}&period_end_date={d}&format=xlsx",
])
def test_reportes_solo_incluyen_el_parqueo(client_a, world, url):
    d = localdate()
    plates = excel_plates(client_a.get(url.format(d=d, m=d.strftime("%Y-%m"))))

    assert A_PLATE + "C" in plates
    assert not any(p.startswith("PBBB") for p in plates)


@pytest.mark.django_db
def test_reporte_por_placa_de_otro_parqueo_no_encuentra_nada(client_a, world):
    d = localdate()
    response = client_a.get(
        f"/parking/reporte/placa/?plate={B_PLATE}C&start_date={d}&end_date={d}&format=xlsx"
    )
    assert response.status_code == 302  # «No se encontraron registros»


# ---------------------------------------------------------------- Pagos de mensualidad

@pytest.mark.django_db
def test_alerta_de_pago_solo_del_parqueo(client_a, world):
    messages = [str(m) for m in client_a.get("/").context["messages"]]
    assert len([m for m in messages if "Pago" in m]) == 1
    assert Payment.objects.filter(subscription__parking=world["B"]).count() == 1


# ---------------------------------------------------------------- Módulo de baños y suspensión

@pytest.mark.django_db
def test_banos_desactivados(client_a, world):
    Parking.objects.filter(pk=world["A"].pk).update(bathrooms_enabled=False)

    assert client_a.get("/baños/entradas/").status_code == 404
    assert client_a.get("/baños/tarifas/").status_code == 404
    assert "Baños" not in body(client_a.get("/")).split("Acciones rápidas")[1]


@pytest.mark.django_db
def test_parqueo_suspendido_cierra_la_sesion(client_a, world):
    Parking.objects.filter(pk=world["A"].pk).update(active=False)

    response = client_a.get("/parking/historial/")

    assert response.status_code == 302 and response.url == "/login/"
    assert client_a.get("/").status_code == 302  # ya no tiene sesión


@pytest.mark.django_db
def test_usuario_sin_parqueo_no_entra(client, django_user_model, parking):
    client.force_login(django_user_model.objects.create_user("huerfano", password="x"))
    response = client.get("/")
    assert response.status_code == 302 and response.url == "/login/"


# ---------------------------------------------------------------- Superusuario

@pytest.mark.django_db
def test_superusuario_cambia_de_parqueo(admin_client, world):
    assert admin_client.get("/").context["current_parking"] == world["A"]

    admin_client.post("/cuenta/parqueo/", {"parking_id": world["B"].pk, "next": "/"})

    context = admin_client.get("/").context
    assert context["current_parking"] == world["B"]
    assert context["total_daily_income"] == Decimal("7.00")


@pytest.mark.django_db
def test_solo_superusuario_cambia_de_parqueo(client_a, world):
    response = client_a.post("/cuenta/parqueo/", {"parking_id": world["B"].pk})
    assert response.status_code == 403


# ---------------------------------------------------------------- Alta de cuentas

@pytest.mark.django_db
def test_crear_cuenta_de_parqueo(client, db):
    parking, owner = create_parking_account(
        owner_username="dueno_sur", owner_password="Clave-segura-123", name="Parqueo Sur", phone="2222-1111",
    )

    assert owner.parking_membership.parking == parking
    assert owner.groups.filter(name=OWNER_GROUP).exists()
    fee = Fee.objects.get(parking=parking)
    assert fee.default and fee.ranges.count() == 11
    assert fee.calculate_fee(30) == Decimal("1.00") and fee.calculate_fee(400) == Decimal("6.00")

    client.force_login(owner)
    response = client.get("/")
    assert response.status_code == 200
    assert response.context["total_daily_income"] == 0


@pytest.mark.django_db
def test_crear_cuenta_desde_el_admin(admin_client):
    response = admin_client.post("/admin/tenants/parking/add/", {
        "name": "Parqueo Este", "address": "Calle 5", "phone": "2222-3333", "capacity": "40",
        "bathrooms_enabled": "on", "active": "on",
        "owner_username": "dueno_este", "owner_password": "Clave-segura-123", "with_default_fee": "on",
    })

    assert response.status_code == 302
    parking = Parking.objects.get(name="Parqueo Este")
    assert parking.members.get().user.username == "dueno_este"
    assert Fee.objects.filter(parking=parking, default=True).exists()


@pytest.mark.django_db
def test_roles_tienen_permisos_esperados(db):
    groups = ensure_role_groups()
    operator = set(groups[OPERATOR_GROUP].permissions.values_list("codename", flat=True))
    owner = set(groups[OWNER_GROUP].permissions.values_list("codename", flat=True))

    assert "add_entry" in operator and "view_statistics_entry" not in operator
    assert "view_statistics_entry" in owner and "delete_entry" not in owner
