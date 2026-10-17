from decimal import Decimal

import pytest
from django.db.models import ProtectedError
from django.utils.timezone import localtime

from parking.models import Entry


@pytest.mark.django_db
class TestFeeProtect:
    def test_no_se_puede_borrar_tarifa_con_entradas(self, fee, make_entry):
        make_entry(duration=30)
        with pytest.raises(ProtectedError):
            fee.delete()
        assert Entry.objects.count() == 1

    def test_se_puede_borrar_tarifa_sin_uso(self, fee):
        fee.delete()


@pytest.mark.django_db
class TestCalculateFee:
    @pytest.mark.parametrize("minute,expected", [
        (1, Decimal("1.00")), (59, Decimal("1.00")), (60, Decimal("2.00")), (500, Decimal("3.00")),
    ])
    def test_cobro_por_rangos(self, fee, minute, expected):
        assert fee.calculate_fee(minute) == expected

    def test_tiempo_menor_al_primer_rango_no_cobra(self, fee):
        fee.ranges.filter(start_minute=0).update(start_minute=15)
        assert fee.calculate_fee(5) == 0

    def test_tarifa_sin_rangos_no_cobra(self, fee):
        fee.ranges.all().delete()
        assert fee.calculate_fee(30) == 0


@pytest.mark.django_db
class TestDeparture:
    def test_registra_salida_y_congela_monto(self, admin_client, make_entry):
        entry = make_entry(minutes_ago=90)
        response = admin_client.post(f"/parking/salida/{entry.pk}")
        entry.refresh_from_db()
        assert response.status_code == 302
        assert entry.state is False
        assert entry.departure_date_hour is not None
        assert entry.final_minutes in (90, 91)  # minutos parciales se redondean hacia arriba
        assert entry.final_amount == Decimal("2.00")

    def test_segunda_salida_no_modifica_la_entrada(self, admin_client, make_entry):
        entry = make_entry(minutes_ago=300, duration=30)
        before = (entry.departure_date_hour, entry.final_minutes, entry.final_amount)

        response = admin_client.post(f"/parking/salida/{entry.pk}")
        entry.refresh_from_db()

        assert response.status_code == 302
        assert (entry.departure_date_hour, entry.final_minutes, entry.final_amount) == before

    def test_get_de_entrada_cerrada_redirige(self, admin_client, make_entry):
        entry = make_entry(duration=30)
        response = admin_client.get(f"/parking/salida/{entry.pk}")
        assert response.status_code == 302

    def test_salida_mensual_no_cobra(self, admin_client, make_entry, monthly_policy):
        entry = make_entry(plate=monthly_policy.plate, minutes_ago=300, fee=None)
        admin_client.post(f"/parking/salida/{entry.pk}")
        entry.refresh_from_db()
        assert entry.final_amount == 0


@pytest.mark.django_db
class TestReports:
    def summary(self, report, title):
        return next(s for s in report["summary"] if s["title"] == title)

    def test_mensualidad_se_suma_una_vez_por_placa(self, make_entry, monthly_policy):
        from parking.services.report_service import generate_month_report
        from parking.models import PlatePolicy

        PlatePolicy.objects.create(parking=monthly_policy.parking, plate="P100002", billing_type="MONTHLY", amount=Decimal("40.00"))
        for days in range(1, 6):
            make_entry(plate=monthly_policy.plate, minutes_ago=60 * 24 * days, duration=480, fee=None)

        report = generate_month_report(PARKING(), now_local())

        assert self.summary(report, "Suscripción mensual")["income"] == Decimal("80.00")

    def test_mensualidad_sin_monto_no_rompe_el_reporte(self, make_entry, parking):
        from parking.services.report_service import generate_month_report
        from parking.models import PlatePolicy

        PlatePolicy.objects.create(parking=parking, plate="P100003", billing_type="MONTHLY", amount=None)
        make_entry(plate="P100003", duration=60, fee=None)

        report = generate_month_report(PARKING(), now_local())

        assert self.summary(report, "Suscripción mensual")["income"] == 0

    def test_entrada_legacy_se_calcula_al_vuelo_sin_guardarse(self, make_entry):
        from parking.services.report_service import generate_day_report

        entry = make_entry(minutes_ago=200, duration=90)
        Entry.objects.filter(pk=entry.pk).update(final_minutes=None, final_amount=None)
        day = localtime(entry.departure_date_hour).date()

        assert generate_day_report(entry.parking, day)["total_income"] == Decimal("2.00")

        entry.refresh_from_db()
        assert entry.final_amount is None

    def test_suscripcion_por_hora_se_reporta_como_tarifa(self, admin_client, make_entry, parking):
        from parking.models import PlatePolicy

        PlatePolicy.objects.create(parking=parking, plate="P300001", billing_type="HOURLY", amount=Decimal("1.00"))
        entry = make_entry(plate="P300001", minutes_ago=200, duration=90)
        day = localtime(entry.departure_date_hour).date().isoformat()

        excel = admin_client.get(f"/parking/reporte/hoy/?date={day}&format=xlsx")
        plate = admin_client.get(
            f"/parking/reporte/placa/?plate=P300001&start_date={day}&end_date={day}&format=xlsx"
        )

        assert excel.status_code == 200
        assert plate.status_code == 200

    def test_entrada_activa_no_suma_ingresos(self, make_entry):
        from parking.services.report_service import generate_day_report

        make_entry(minutes_ago=30)

        assert generate_day_report(PARKING(), now_local())["total_income"] == 0


def PARKING():
    from tenants.models import Parking
    return Parking.objects.order_by("pk").first()


def now_local():
    from django.utils.timezone import localdate
    return localdate()


@pytest.mark.django_db(transaction=True)
class TestUniqueActiveEntry:
    def test_base_de_datos_impide_dos_entradas_activas(self, fee):
        from django.db import IntegrityError

        Entry.objects.bulk_create([Entry(parking=fee.parking, plate="P555555", fee=fee)])
        with pytest.raises(IntegrityError):
            Entry.objects.bulk_create([Entry(parking=fee.parking, plate="P555555", fee=fee)])

    def test_save_traduce_el_error_a_mensaje_claro(self, make_entry):
        make_entry(plate="P555555")
        with pytest.raises(ValueError, match="Ya existe una entrada activa"):
            make_entry(plate="P555555")

    def test_misma_placa_puede_tener_varias_entradas_cerradas(self, make_entry):
        make_entry(plate="P555555", minutes_ago=300, duration=30)
        make_entry(plate="P555555", minutes_ago=200, duration=30)
        make_entry(plate="P555555")
        assert Entry.objects.filter(plate="P555555").count() == 3


@pytest.mark.django_db
def test_excel_deja_vacia_la_salida_de_entradas_activas(admin_client, make_entry):
    from io import BytesIO
    from openpyxl import load_workbook

    make_entry(plate="PACTIVA", minutes_ago=30)
    make_entry(plate="PCERRADA", minutes_ago=200, duration=60)

    response = admin_client.get(f"/parking/reporte/hoy/?date={now_local()}&format=xlsx")
    rows = {r[0]: r for r in load_workbook(BytesIO(response.content)).active.iter_rows(min_row=4, values_only=True)}

    assert rows["PACTIVA"][2] in ("", None)
    assert rows["PCERRADA"][2]


@pytest.mark.django_db
class TestTicket:
    @pytest.mark.parametrize("query", ["", "?entry_id=", "?entry_id=abc", "?entry_id=999999"])
    def test_ticket_invalido_da_404(self, admin_client, query):
        assert admin_client.get(f"/parking/ticket/{query}").status_code == 404

    def test_ticket_valido(self, admin_client, make_entry):
        entry = make_entry()
        assert admin_client.get(f"/parking/ticket/?entry_id={entry.pk}").status_code == 200

    def test_ticket_requiere_permiso(self, client, django_user_model, make_entry, member):
        client.force_login(member(django_user_model.objects.create_user("sin_permiso", password="x")))
        assert client.get(f"/parking/ticket/?entry_id={make_entry().pk}").status_code == 403


class TestPdfStaticFiles:
    def test_css_se_lee_de_disco_sin_peticion_http(self):
        from pathlib import Path
        from django.conf import settings
        from parking.utils import StaticFilesURLFetcher

        fetcher = StaticFilesURLFetcher("http://servidor-inexistente.invalid/")
        response = fetcher.fetch("http://servidor-inexistente.invalid/static/css/styles.css")
        body = response.read()

        assert body == (Path(settings.BASE_DIR) / "static/css/styles.css").read_bytes()
        assert response.headers.get_content_type() == "text/css"

    def test_estatico_inexistente_falla_sin_ir_a_la_red(self):
        from parking.utils import StaticFilesURLFetcher

        fetcher = StaticFilesURLFetcher("http://servidor-inexistente.invalid/")
        with pytest.raises(FileNotFoundError):
            fetcher.fetch("http://servidor-inexistente.invalid/static/no-existe.css")


@pytest.mark.django_db
class TestConsultasConstantes:
    """El número de consultas no debe crecer con la cantidad de registros (N+1)."""

    def count_queries(self, client, url):
        from django.db import connection
        from django.test.utils import CaptureQueriesContext

        with CaptureQueriesContext(connection) as ctx:
            assert client.get(url).status_code == 200
        return len(ctx.captured_queries)

    def add_entries(self, make_entry, prefix, n):
        for i in range(n):
            make_entry(plate=f"{prefix}{i}", minutes_ago=300 - i, duration=60 if i % 2 else None)

    @pytest.mark.parametrize("url", [
        "/parking/historial/",
        "/parking/reporte/hoy/?date={d}&format=xlsx",
    ])
    def test_parqueo(self, admin_client, make_entry, url):
        url = url.format(d=now_local())
        self.add_entries(make_entry, "PA", 2)
        few = self.count_queries(admin_client, url)

        self.add_entries(make_entry, "PB", 10)
        many = self.count_queries(admin_client, url)

        assert many == few

    def test_historial_muestra_lo_cobrado_aunque_cambie_la_tarifa(self, admin_client, make_entry, fee):
        entry = make_entry(minutes_ago=200, duration=90)
        fee.ranges.update(amount=Decimal("99.00"))

        response = admin_client.get("/parking/historial/")
        shown = next(e for e in response.context["entries"] if e.pk == entry.pk)

        assert shown.amount == Decimal("2.00")


@pytest.mark.django_db
def test_activar_y_desactivar_suscripcion(admin_client, monthly_policy):
    url = f"/parking/suscripciones/estado/{monthly_policy.pk}"

    admin_client.post(url)
    monthly_policy.refresh_from_db()
    assert monthly_policy.active is False

    admin_client.post(url)
    monthly_policy.refresh_from_db()
    assert monthly_policy.active is True


class TestDictadoPlaca:
    @pytest.mark.parametrize("dicho,placa", [
        # Palabras agrupadas como se lee la placa
        ("pe ciento setenta y uno setecientos treinta y siete", "P171737"),
        ("pe cuarenta ochocientos siete", "P40807"),
        ("eme novecientos once ciento dieciséis", "M911116"),
        ("ce veintidós cero cero cinco", "C22005"),
        # Dígito por dígito
        ("pe uno siete uno siete tres siete", "P171737"),
        ("pe cero cero uno", "P001"),
        # El navegador ya entrega dígitos o letras
        ("P 171 737", "P171737"),
        ("p171737", "P171737"),
        ("P-40807", "P40807"),
        ("pe 171737", "P171737"),
        # Placas alfanuméricas y prefijos de dos letras
        ("pe ocho e noventa y ocho", "P8E98"),
        ("a be cien", "AB100"),
        ("a b 100", "AB100"),
        # Relleno, acentos y mayúsculas
        ("la placa es pe cuatro cero ocho cero siete", "P40807"),
        ("Placa número PE 40 807", "P40807"),
        ("pe cuatro mil ochocientos siete", "P4807"),
        # Letra al final
        ("P11017E", "P11017E"),
        ("P 11017 e", "P11017E"),
        ("pe once mil diecisiete e", "P11017E"),
        ("pe once cero diecisiete letra e", "P11017E"),
        ("P 11017 letra de", "P11017D"),
        ("P 11017 i griega", "P11017Y"),
        ("P 11017 doble ve", "P11017W"),
    ])
    def test_interpreta_lo_dicho(self, dicho, placa):
        from parking.services.voice_plate import parse_plate
        assert parse_plate(dicho) == placa

    @pytest.mark.parametrize("dicho", ["", "hola", "ciento setenta", "pe"])
    def test_texto_sin_placa_valida(self, dicho):
        from parking.services.voice_plate import is_valid_plate, parse_plate
        assert not is_valid_plate(parse_plate(dicho))


@pytest.mark.django_db
class TestDictadoEndpoint:
    url = "/parking/busqueda/voz/"

    def post(self, client, alternatives, prefix=None):
        import json
        payload = {"alternatives": alternatives}
        if prefix is not None:
            payload["prefix"] = prefix
        return client.post(self.url, json.dumps(payload), content_type="application/json")

    def test_placa_nueva_registrara_entrada(self, admin_client, fee):
        data = self.post(admin_client, ["pe ciento setenta y uno setecientos treinta y siete"]).json()

        assert data["ok"] is True
        assert data["plate"] == "P171737"
        assert data["formatted"] == "P 171 737"
        assert data["action"] == "entrada"

    def test_placa_adentro_registrara_salida(self, admin_client, make_entry):
        make_entry(plate="P171737")
        data = self.post(admin_client, ["P 171 737"]).json()

        assert data["action"] == "salida"
        assert data["entry_time"]

    def test_sugiere_placa_parecida_que_esta_adentro(self, admin_client, make_entry):
        make_entry(plate="P171737")
        data = self.post(admin_client, ["P 171 787"]).json()

        assert data["plate"] == "P171787"
        assert data["candidates"][0]["plate"] == "P171737"
        assert data["candidates"][0]["active"] is True

    def test_prefiere_la_alternativa_que_es_placa_conocida(self, admin_client, make_entry):
        make_entry(plate="P171737")
        data = self.post(admin_client, ["pe ciento setenta y uno setecientos ochenta y siete", "P 171 737"]).json()

        assert data["plate"] == "P171737"
        assert data["action"] == "salida"

    def test_texto_sin_placa(self, admin_client):
        data = self.post(admin_client, ["buenos días"]).json()
        assert data["ok"] is False

    def test_ofrece_la_lectura_con_letra_final(self, admin_client, fee):
        # La primera lectura omitió la «e»; otra alternativa sí la trae
        data = self.post(admin_client, ["P 11017", "P 11017 e"]).json()

        assert data["plate"] == "P11017"
        assert "P11017E" in [c["plate"] for c in data["candidates"]]

    @pytest.mark.parametrize("dicho", ["e", "letra e", "E"])
    def test_segundo_dictado_agrega_la_letra_que_falto(self, admin_client, fee, dicho):
        data = self.post(admin_client, [dicho], prefix="P11017").json()

        assert data["ok"] is True
        assert data["plate"] == "P11017E"

    def test_segundo_dictado_con_placa_completa_la_reemplaza(self, admin_client, fee):
        data = self.post(admin_client, ["P 40 807"], prefix="P11017").json()
        assert data["plate"] == "P40807"

    def test_prefijo_invalido_se_ignora(self, admin_client, fee):
        data = self.post(admin_client, ["e"], prefix="hola").json()
        assert data["ok"] is False

    @pytest.mark.parametrize("body", ["no es json", '{"alternatives": "texto"}'])
    def test_solicitud_invalida(self, admin_client, body):
        response = admin_client.post(self.url, body, content_type="application/json")
        assert response.status_code == 400

    def test_requiere_permiso(self, client, django_user_model, member):
        client.force_login(member(django_user_model.objects.create_user("sin_permiso", password="x")))
        assert self.post(client, ["P 171 737"]).status_code == 403

    def test_requiere_csrf(self, django_user_model, parking):
        import json
        from django.test import Client

        client = Client(enforce_csrf_checks=True)
        client.force_login(django_user_model.objects.create_superuser("a", "a@a.com", "x"))
        response = client.post(self.url, json.dumps({"alternatives": ["P1"]}), content_type="application/json")

        assert response.status_code == 403

    def test_pantalla_de_busqueda_tiene_boton_y_modal(self, admin_client):
        html = admin_client.get("/parking/busqueda/").content.decode()

        assert "Escanear ticket" in html
        assert "Dictar placa" in html
        assert 'id="voicePlateModal"' in html
        assert 'id="plateInput"' in html


@pytest.mark.django_db
class TestEliminarEntrada:
    def url(self, entry):
        return f"/parking/eliminar/{entry.pk}/"

    def test_superusuario_elimina_y_queda_registro(self, admin_client, make_entry):
        from django.contrib.admin.models import DELETION, LogEntry

        entry = make_entry(plate="P171737", minutes_ago=120, duration=60)
        response = admin_client.post(self.url(entry))

        assert response.status_code == 302
        assert response.url == "/parking/historial/"
        assert not Entry.objects.filter(pk=entry.pk).exists()

        log = LogEntry.objects.get(action_flag=DELETION)
        assert log.user.username == "admin"
        assert log.object_id == str(entry.pk)
        assert "P171737" in log.change_message
        assert "$2.00" in log.change_message

    def test_usuario_con_permiso_elimina(self, client, django_user_model, make_entry, member):
        from django.contrib.auth.models import Permission

        user = member(django_user_model.objects.create_user("jefe", password="x"))
        user.user_permissions.add(Permission.objects.get(codename="delete_entry"))
        client.force_login(user)
        entry = make_entry()

        assert client.post(self.url(entry)).status_code == 302
        assert not Entry.objects.filter(pk=entry.pk).exists()

    def test_permiso_por_grupo_tambien_sirve(self, client, django_user_model, make_entry, member):
        from django.contrib.auth.models import Group, Permission

        group = Group.objects.create(name="Administrador")
        group.permissions.add(Permission.objects.get(codename="delete_entry"))
        user = member(django_user_model.objects.create_user("jefe", password="x"))
        user.groups.add(group)
        client.force_login(user)
        entry = make_entry()

        assert client.post(self.url(entry)).status_code == 302
        assert not Entry.objects.filter(pk=entry.pk).exists()

    def test_sin_permiso_no_puede(self, client, django_user_model, make_entry, member):
        from django.contrib.auth.models import Permission

        user = member(django_user_model.objects.create_user("editor", password="x", is_staff=True))
        user.user_permissions.set(
            Permission.objects.filter(content_type__app_label="parking").exclude(codename="delete_entry")
        )
        client.force_login(user)
        entry = make_entry()

        assert client.post(self.url(entry)).status_code == 403
        assert Entry.objects.filter(pk=entry.pk).exists()

    def test_get_no_elimina(self, admin_client, make_entry):
        entry = make_entry()
        assert admin_client.get(self.url(entry)).status_code == 405
        assert Entry.objects.filter(pk=entry.pk).exists()

    def test_anonimo_va_al_login(self, client, make_entry):
        entry = make_entry()
        response = client.post(self.url(entry))
        assert response.status_code == 302 and response.url.startswith("/login/")
        assert Entry.objects.filter(pk=entry.pk).exists()

    def test_entrada_inexistente(self, admin_client):
        assert admin_client.post("/parking/eliminar/999999/").status_code == 404

    def test_boton_solo_visible_con_permiso(self, admin_client, client, django_user_model, make_entry, member):
        from django.contrib.auth.models import Permission

        entry = make_entry()
        assert "Eliminar entrada" in admin_client.get(f"/parking/editar/{entry.pk}/").content.decode()

        user = member(django_user_model.objects.create_user("editor", password="x"))
        user.user_permissions.add(Permission.objects.get(codename="change_entry"))
        client.force_login(user)
        html = client.get(f"/parking/editar/{entry.pk}/").content.decode()
        assert "Eliminar entrada" not in html
        assert "deleteEntryModal" not in html

        user.user_permissions.add(Permission.objects.get(codename="delete_entry"))
        user = django_user_model.objects.get(pk=user.pk)  # limpia la caché de permisos
        client.force_login(user)
        assert "deleteEntryModal" in client.get(f"/parking/editar/{entry.pk}/").content.decode()


@pytest.mark.django_db
class TestEscaneoPlaca:
    url = "/parking/busqueda/foto/"

    @pytest.fixture(autouse=True)
    def setup(self, settings, parking, monkeypatch):
        from types import SimpleNamespace
        from django.core.cache import cache
        from parking.services import plate_scan

        from payment.models import Plan, Subscription
        from tenants.models import AiPackage

        settings.ANTHROPIC_API_KEY = "test-key"
        # El parqueo tiene un plan con IA incluida
        self.package = AiPackage.objects.create(name="Pro", scans=1000, price=Decimal("12"))
        self.plan = Plan.objects.create(name="Pro", price=Decimal("45"), ai_package=self.package)
        self.subscription = Subscription.objects.create(parking=parking, plan=self.plan, active=True)
        cache.clear()

        self.calls = []
        self.reply = "P11017E"

        def create(**kwargs):
            self.calls.append(kwargs)
            if isinstance(self.reply, Exception):
                raise self.reply
            return SimpleNamespace(
                content=[SimpleNamespace(type="text", text=self.reply)],
                usage=SimpleNamespace(input_tokens=400, output_tokens=8),
            )

        self.keys = []

        def client(api_key):
            self.keys.append(api_key)
            return SimpleNamespace(messages=SimpleNamespace(create=create))

        monkeypatch.setattr(plate_scan, "_client", client)

    @staticmethod
    def image(width=600, height=300):
        import io
        from PIL import Image

        out = io.BytesIO()
        Image.new("RGB", (width, height), "white").save(out, format="JPEG")
        out.seek(0)
        out.name = "placa.jpg"
        return out

    def post(self, client, image=None):
        return client.post(self.url, {"image": image or self.image()})

    def test_lee_la_placa_y_registra_el_costo(self, admin_client, parking, fee):
        from parking.models import PlateScan

        data = self.post(admin_client).json()

        assert data["ok"] is True
        assert data["plate"] == "P11017E"
        assert data["action"] == "entrada"

        scan = PlateScan.objects.get()
        assert scan.parking == parking
        assert scan.user.username == "admin"
        assert (scan.input_tokens, scan.output_tokens) == (400, 8)
        assert scan.cost == Decimal("0.000440")  # 400 × $1/M + 8 × $5/M

    def test_placa_adentro_registrara_salida(self, admin_client, make_entry):
        make_entry(plate="P11017E")
        assert self.post(admin_client).json()["action"] == "salida"

    def test_lectura_dudosa_prefiere_la_placa_conocida(self, admin_client, make_entry):
        make_entry(plate="P11017E")
        self.reply = "P11O17E, P11017E"

        data = self.post(admin_client).json()

        assert data["plate"] == "P11017E"
        assert data["action"] == "salida"

    def test_sin_placa_legible(self, admin_client, fee):
        from parking.models import PlateScan

        self.reply = "NINGUNA"
        data = self.post(admin_client).json()

        assert data["ok"] is False
        assert PlateScan.objects.count() == 1  # la lectura igual tuvo costo

    def test_imagen_grande_se_reduce_antes_de_enviar(self, admin_client, fee):
        import base64
        import io
        from PIL import Image

        self.post(admin_client, self.image(3000, 1500))

        sent = self.calls[0]["messages"][0]["content"][0]["source"]["data"]
        assert max(Image.open(io.BytesIO(base64.b64decode(sent))).size) <= 800

    def test_archivo_que_no_es_imagen(self, admin_client):
        from django.core.files.uploadedfile import SimpleUploadedFile

        response = admin_client.post(self.url, {"image": SimpleUploadedFile("x.jpg", b"no es imagen")})

        assert response.status_code == 400
        assert self.calls == []

    def test_sin_imagen(self, admin_client):
        assert admin_client.post(self.url).status_code == 400

    def test_error_de_la_api(self, admin_client):
        from parking.models import PlateScan

        self.reply = RuntimeError("timeout")
        response = self.post(admin_client)

        assert response.status_code == 400
        assert "Intenta de nuevo" in response.json()["error"]
        assert not PlateScan.objects.exists()

    def test_parqueo_sin_el_servicio(self, admin_client, parking):
        self.subscription.active = False
        self.subscription.save()

        assert self.post(admin_client).status_code == 403
        assert self.calls == []

    def test_plan_sin_ia(self, admin_client, parking):
        self.plan.ai_package = None
        self.plan.save()

        assert self.post(admin_client).status_code == 403

    def test_plan_sin_ia_con_recarga_del_mes(self, admin_client, parking, fee):
        from tenants.models import AiTopUp

        self.plan.ai_package = None
        self.plan.save()
        AiTopUp.objects.create(parking=parking, package=self.package)

        assert self.post(admin_client).status_code == 200

    def test_ia_de_todos_los_planes_activos_se_suma(self, admin_client, parking, fee):
        from payment.models import Plan, Subscription
        from tenants.models import AiPackage
        from parking.services import plate_scan

        addon = AiPackage.objects.create(name="Complemento", scans=500, price=Decimal("6"))
        Subscription.objects.create(
            parking=parking, plan=Plan.objects.create(name="IA extra", price=Decimal("6"), ai_package=addon),
            active=True,
        )
        assert plate_scan.quota(parking)["limit"] == 1500

    def test_sin_api_key(self, admin_client, settings):
        settings.ANTHROPIC_API_KEY = ""
        assert self.post(admin_client).status_code == 403

    def give_package(self, parking, scans):
        self.package.scans = scans
        self.package.save()
        return self.package

    def test_paquete_agotado_bloquea_el_escaneo(self, admin_client, parking, fee):
        self.give_package(parking, 2)

        assert self.post(admin_client).status_code == 200
        assert self.post(admin_client).status_code == 200
        response = self.post(admin_client)

        assert response.status_code == 429
        assert "paquete de IA" in response.json()["error"]
        assert len(self.calls) == 2

    def test_paquete_extra_del_mes_amplia_el_limite(self, admin_client, parking, fee):
        from tenants.models import AiTopUp

        package = self.give_package(parking, 1)
        assert self.post(admin_client).status_code == 200
        assert self.post(admin_client).status_code == 429

        top_up = AiTopUp.objects.create(parking=parking, package=package)
        assert (top_up.scans, top_up.price) == (1, Decimal("12"))
        assert self.post(admin_client).status_code == 200
        assert self.post(admin_client).status_code == 429

    def test_paquete_extra_de_otro_mes_no_cuenta(self, admin_client, parking, fee):
        from datetime import date
        from tenants.models import AiTopUp

        package = self.give_package(parking, 1)
        AiTopUp.objects.create(parking=parking, package=package, month=date(2020, 1, 1))
        assert self.post(admin_client).status_code == 200
        assert self.post(admin_client).status_code == 429

    def test_escaneos_del_mes_pasado_no_cuentan(self, admin_client, parking, fee):
        from datetime import timedelta
        from django.utils.timezone import now
        from parking.models import PlateScan

        self.give_package(parking, 1)
        PlateScan.objects.create(parking=parking, model="x", created_at=now() - timedelta(days=40))
        assert self.post(admin_client).status_code == 200

    def test_limite_por_minuto(self, admin_client, fee):
        from parking.views import SCAN_RATE_PER_MINUTE

        for _ in range(SCAN_RATE_PER_MINUTE):
            assert self.post(admin_client).status_code == 200

        assert self.post(admin_client).status_code == 429

    def test_el_limite_cuenta_solo_el_parqueo_propio(self, admin_client, parking, other_parking, fee):
        from parking.models import PlateScan

        PlateScan.objects.create(parking=other_parking, model="x")
        self.give_package(parking, 1)

        assert self.post(admin_client).status_code == 200

    def test_usa_la_api_key_y_el_modelo_del_admin(self, admin_client, fee):
        from parking.models import PlateScan
        from tenants.models import AiSettings

        config = AiSettings.load()
        config.api_key = "sk-ant-admin"
        config.model = "claude-sonnet-5"
        config.price_input = Decimal("3")
        config.price_output = Decimal("15")
        config.save()

        self.post(admin_client)

        assert self.keys == ["sk-ant-admin"]
        assert self.calls[0]["model"] == "claude-sonnet-5"
        assert PlateScan.objects.get().cost == Decimal("0.001320")  # 400 × $3/M + 8 × $15/M

    def test_boton_desactivado_y_alerta_cuando_se_agota(self, admin_client, parking, fee):
        from parking.models import PlateScan

        self.give_package(parking, 1)
        PlateScan.objects.create(parking=parking, model="x")

        html = admin_client.get("/parking/busqueda/").content.decode()
        assert "Paquete de IA agotado" in html
        assert "plateCameraModal" not in html
        assert "Tu paquete de IA para escanear placas se terminó" in html

        # El resto sigue funcionando
        assert admin_client.get("/").status_code == 200

    def test_alerta_cuando_queda_poco(self, admin_client, parking, fee):
        from parking.models import PlateScan

        self.give_package(parking, 10)
        PlateScan.objects.bulk_create(PlateScan(parking=parking, model="x") for _ in range(8))

        html = admin_client.get("/parking/busqueda/").content.decode()
        assert "Te quedan 2 de 10 escaneos" in html
        assert "Escanear placa" in html

    def test_operador_no_ve_la_alerta_pero_si_el_boton_desactivado(self, client, django_user_model, member, parking, fee):
        from django.contrib.auth.models import Permission
        from parking.models import PlateScan

        self.give_package(parking, 1)
        PlateScan.objects.create(parking=parking, model="x")
        user = member(django_user_model.objects.create_user("op", password="x"))
        user.user_permissions.add(Permission.objects.get(codename="add_entry"))
        client.force_login(user)

        html = client.get("/parking/busqueda/").content.decode()
        assert "Paquete de IA agotado" in html
        assert "aiScanAlert" not in html

    def test_requiere_permiso(self, client, django_user_model, member):
        client.force_login(member(django_user_model.objects.create_user("sin_permiso", password="x")))
        assert self.post(client).status_code == 403

    def test_boton_visible_solo_con_el_servicio(self, admin_client, parking, settings):
        html = admin_client.get("/parking/busqueda/").content.decode()
        assert "Escanear placa" in html and 'id="plateCameraModal"' in html

        self.subscription.active = False
        self.subscription.save()
        html = admin_client.get("/parking/busqueda/").content.decode()
        assert "plateCameraModal" not in html

        self.subscription.active = True
        self.subscription.save()
        settings.ANTHROPIC_API_KEY = ""
        assert "plateCameraModal" not in admin_client.get("/parking/busqueda/").content.decode()

    def test_admin_muestra_totales(self, admin_client, parking):
        from parking.models import PlateScan

        PlateScan.objects.create(parking=parking, model="x", cost=Decimal("0.0005"))
        PlateScan.objects.create(parking=parking, model="x", cost=Decimal("0.0007"))

        html = admin_client.get("/admin/parking/platescan/").content.decode()
        assert "<strong>2</strong> escaneos" in html
        assert "$0.0012" in html


@pytest.mark.parametrize("texto,lecturas", [
    ("P11017E", ["P11017E"]),
    ("P 110-17E", ["P11017E"]),
    ("P11O17E, P11017E", ["P11O17E", "P11017E"]),
    ("NINGUNA", []),
    ("", []),
])
def test_lecturas_de_la_ia(texto, lecturas):
    from parking.services.plate_scan import parse_readings
    assert parse_readings(texto) == lecturas


@pytest.mark.django_db
class TestAdminIA:
    def test_api_key_se_guarda_oculta_y_se_conserva(self, admin_client):
        from tenants.models import AiSettings

        url = "/admin/tenants/aisettings/1/change/"
        data = {"model": "claude-haiku-4-5-20251001", "price_input": "1", "price_output": "5"}

        assert admin_client.get("/admin/tenants/aisettings/").url == url

        admin_client.post(url, {**data, "api_key": "sk-ant-secreta-1234"})
        assert AiSettings.load().api_key == "sk-ant-secreta-1234"

        html = admin_client.get(url).content.decode()
        assert "sk-ant-secreta-1234" not in html
        assert "…1234" in html

        admin_client.post(url, {**data, "api_key": ""})
        assert AiSettings.load().api_key == "sk-ant-secreta-1234"

        admin_client.post(url, {**data, "api_key": "", "clear_api_key": "on"})
        assert AiSettings.load().api_key == ""

    def test_rechaza_una_clave_con_formato_invalido(self, admin_client):
        from tenants.models import AiSettings

        admin_client.post("/admin/tenants/aisettings/1/change/", {
            "model": "x", "price_input": "1", "price_output": "5", "api_key": "hola",
        })
        assert AiSettings.load().api_key == ""

    def test_paquetes_muestran_costo_y_ganancia(self, admin_client):
        from tenants.models import AiPackage

        AiPackage.objects.create(name="Básico", scans=3000, price=Decimal("5"))
        html = admin_client.get("/admin/tenants/aipackage/").content.decode()

        assert "$0.90" in html  # 3,000 × $0.0003 de referencia
        assert "$4.10 (82%)" in html

    def test_lista_de_parqueos_muestra_plan_y_uso_de_ia(self, admin_client, parking):
        from parking.models import PlateScan
        from payment.models import Plan, Subscription
        from tenants.models import AiPackage

        package = AiPackage.objects.create(name="IA 3,000", scans=3000, price=Decimal("12"))
        plan = Plan.objects.create(name="Pro", price=Decimal("45"), ai_package=package)
        Subscription.objects.create(parking=parking, plan=plan, active=True)
        PlateScan.objects.create(parking=parking, model="x")

        html = admin_client.get("/admin/tenants/parking/").content.decode()
        assert "1 / 3,000" in html
        assert ">Pro<" in html


@pytest.mark.django_db
def test_reporte_por_periodo_ya_no_existe(admin_client):
    # Se quitó: un PDF de muchos meses agotaba la memoria del servidor
    html = admin_client.get("/parking/reportes/").content.decode()

    assert "tab-period" not in html
    assert "Mensual" in html and "Por placa" in html
    assert admin_client.get("/parking/reporte/periodo/").status_code == 404


# ---------------------------------------------------------------- Llaves en custodia

@pytest.mark.django_db
class TestLlaves:

    @pytest.fixture(autouse=True)
    def keys_on(self, parking):
        parking.keys_enabled = True
        parking.save()

    @staticmethod
    def photo(width=3000, height=2000):
        import io
        from django.core.files.uploadedfile import SimpleUploadedFile
        from PIL import Image

        out = io.BytesIO()
        Image.new("RGB", (width, height), "orange").save(out, format="JPEG", quality=95)
        return SimpleUploadedFile("llave.jpg", out.getvalue(), content_type="image/jpeg")

    def save(self, client, entry, hook="", photo=True):
        data = {"hook": hook}
        if photo:
            data["photo"] = self.photo()
        return client.post(f"/parking/llaves/{entry.pk}/", data)

    def test_guarda_foto_optimizada_y_gancho(self, admin_client, make_entry):
        import io
        from PIL import Image
        from parking.models import KeyDeposit

        entry = make_entry(plate="P11017E")
        response = self.save(admin_client, entry, hook="7")

        assert response.json() == {"ok": True, "hook": "7"}
        deposit = KeyDeposit.objects.get(entry=entry)
        image = Image.open(io.BytesIO(bytes(deposit.photo)))
        assert max(image.size) <= 1024 and image.format == "JPEG"
        assert deposit.created_by.username == "admin"
        assert deposit.parking == entry.parking

    def test_cambiar_solo_el_gancho_conserva_la_foto(self, admin_client, make_entry):
        from parking.models import KeyDeposit

        entry = make_entry()
        self.save(admin_client, entry, hook="7")
        before = bytes(KeyDeposit.objects.get(entry=entry).photo)

        assert self.save(admin_client, entry, hook="12", photo=False).json()["hook"] == "12"
        deposit = KeyDeposit.objects.get(entry=entry)
        assert bytes(deposit.photo) == before

    def test_la_primera_vez_exige_foto(self, admin_client, make_entry):
        response = self.save(admin_client, make_entry(), hook="3", photo=False)
        assert response.status_code == 400

    def test_archivo_que_no_es_imagen(self, admin_client, make_entry):
        from django.core.files.uploadedfile import SimpleUploadedFile

        response = admin_client.post(f"/parking/llaves/{make_entry().pk}/", {
            "photo": SimpleUploadedFile("x.jpg", b"no es imagen"),
        })
        assert response.status_code == 400

    def test_no_se_asigna_a_un_carro_que_ya_salio(self, admin_client, make_entry):
        entry = make_entry(minutes_ago=120, duration=60)
        assert self.save(admin_client, entry).status_code == 404

    def test_foto_solo_mientras_el_carro_esta_adentro(self, admin_client, make_entry):
        entry = make_entry()
        self.save(admin_client, entry, hook="7")

        response = admin_client.get(f"/parking/llaves/{entry.pk}/foto/")
        assert response.status_code == 200
        assert response["Content-Type"] == "image/jpeg"
        assert "private" in response["Cache-Control"]

    def test_al_cobrar_la_salida_se_borra_la_foto(self, admin_client, make_entry):
        from parking.models import KeyDeposit

        entry = make_entry()
        self.save(admin_client, entry, hook="7")

        admin_client.post(f"/parking/salida/{entry.pk}")

        assert not KeyDeposit.objects.exists()
        assert admin_client.get(f"/parking/llaves/{entry.pk}/foto/").status_code == 404

    def test_al_cerrar_la_entrada_por_cualquier_via_se_borra(self, admin_client, make_entry):
        from django.utils.timezone import now
        from parking.models import KeyDeposit

        entry = make_entry()
        self.save(admin_client, entry)
        entry.departure_date_hour = now()
        entry.save()

        assert not KeyDeposit.objects.exists()

    def test_lista_ordenada_por_gancho_y_solo_carros_adentro(self, admin_client, make_entry):
        for plate, hook in [("PA1", "10"), ("PA2", "2"), ("PA3", ""), ("PA4", "B")]:
            self.save(admin_client, make_entry(plate=plate), hook=hook)

        html = admin_client.get("/parking/llaves/").content.decode()
        order = [html.index(p) for p in ["PA2", "PA1", "PA4", "PA3"]]
        assert order == sorted(order)
        assert "4 llaves" in html

    def test_botones_en_salida_historial_menu_y_dashboard(self, admin_client, make_entry):
        entry = make_entry(plate="P11017E")

        departure = admin_client.get(f"/parking/salida/{entry.pk}").content.decode()
        assert "Agregar llave" in departure and 'id="keyCaptureModal"' in departure

        self.save(admin_client, entry, hook="9")
        departure = admin_client.get(f"/parking/salida/{entry.pk}").content.decode()
        assert "#9" in departure and "keyViewModal" in departure

        record = admin_client.get("/parking/historial/").content.decode()
        assert 'data-key-hook="9"' in record

        dashboard = admin_client.get("/").content.decode()
        assert "1 en custodia" in dashboard
        assert "/parking/llaves/" in record  # menú inferior

    def test_parqueo_sin_el_modulo(self, admin_client, parking, make_entry):
        parking.keys_enabled = False
        parking.save()
        entry = make_entry()

        assert admin_client.get("/parking/llaves/").status_code == 404
        assert self.save(admin_client, entry).status_code == 404
        assert "key-btn" not in admin_client.get(f"/parking/salida/{entry.pk}").content.decode()
        assert "/parking/llaves/" not in admin_client.get("/parking/historial/").content.decode()

    def test_requiere_permiso_para_guardar(self, client, django_user_model, member, make_entry):
        from django.contrib.auth.models import Permission

        user = member(django_user_model.objects.create_user("solo_ver", password="x"))
        user.user_permissions.add(Permission.objects.get(codename="view_entry"))
        client.force_login(user)
        entry = make_entry()

        assert client.get("/parking/llaves/").status_code == 200
        assert self.save(client, entry).status_code == 403

    def test_aislamiento_entre_parqueos(self, client, django_user_model, member, make_entry, other_parking):
        from django.contrib.auth.models import Permission
        from parking.models import Entry

        other_parking.keys_enabled = True
        other_parking.save()
        foreign = Entry.objects.create(parking=other_parking, plate="OTRO1")

        user = member(django_user_model.objects.create_user("op", password="x"))
        user.user_permissions.add(*Permission.objects.filter(codename__in=["view_entry", "add_entry"]))
        client.force_login(user)

        assert self.save(client, foreign).status_code == 404
        assert client.get(f"/parking/llaves/{foreign.pk}/foto/").status_code == 404
